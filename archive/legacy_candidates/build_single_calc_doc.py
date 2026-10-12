# -*- coding: utf-8 -*-
"""
Tự động biên soạn và cập nhật DUY NHẤT 1 file Word tính toán:
D:\HoMinhKhang\PROJECT\GITHUB_PROJECT\CODE XE TU HANH\TINH_TOAN_THIET_KE_XE_TU_HANH.docx

Nội dung giải thích rõ ràng thực tế:
- Tốc độ tính toán lý thuyết cần dùng: n_yeucau = 130 - 140 RPM (để xe đạt v = 0.5 m/s).
- Thực tế phần cứng đã mua: Động cơ GA25-370 280 RPM (n_thucte = 280 RPM).
- Biện pháp giải quyết thực tế: Khống chế xung PWM ở mức 45% - 50% trong code điều khiển vi điều khiển (hoặc gửi qua I2C cho module MKE-M18), giải thích rõ tại sao không bị lãng phí công suất.
"""
import os
import lxml.etree as ET
import latex2mathml.converter
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml

target_dir = r"D:\HoMinhKhang\PROJECT\GITHUB_PROJECT\CODE XE TU HANH"
os.makedirs(target_dir, exist_ok=True)
out_docx = os.path.join(target_dir, "TINH_TOAN_THIET_KE_XE_TU_HANH.docx")

xslt_path = r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"
xslt = ET.parse(xslt_path)
xslt_transform = ET.XSLT(xslt)

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(0.8)

BLACK = RGBColor(0, 0, 0)
PRIMARY = RGBColor(0, 51, 102)

def format_p(p, ls=1.3, sa=4, sb=0, al=WD_ALIGN_PARAGRAPH.LEFT, indent=0):
    p.paragraph_format.line_spacing = ls
    p.paragraph_format.space_after = Pt(sa)
    p.paragraph_format.space_before = Pt(sb)
    p.alignment = al
    if indent > 0:
        p.paragraph_format.left_indent = Inches(indent)
    return p

def add_title(text):
    p = doc.add_paragraph()
    format_p(p, 1.2, 4, 10, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = BLACK

def add_subtitle(text):
    p = doc.add_paragraph()
    format_p(p, 1.2, 14, 2, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = PRIMARY

def add_h1(text):
    p = doc.add_paragraph()
    format_p(p, 1.2, 4, 12, WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = PRIMARY

def add_h2(text):
    p = doc.add_paragraph()
    format_p(p, 1.2, 3, 8, WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = BLACK

def add_body(text, bold_prefix="", indent=0):
    p = doc.add_paragraph()
    format_p(p, 1.3, 3, 0, WD_ALIGN_PARAGRAPH.LEFT, indent)
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.font.name = "Times New Roman"
        rb.font.size = Pt(12)
        rb.font.bold = True
        rb.font.color.rgb = BLACK
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.color.rgb = BLACK

def add_eq(latex_code):
    mathml = latex2mathml.converter.convert(latex_code)
    tree = ET.fromstring(mathml)
    omml = xslt_transform(tree)
    ostr = ET.tostring(omml, encoding="utf-8").decode("utf-8")
    p = doc.add_paragraph()
    format_p(p, 1.15, 4, 4, WD_ALIGN_PARAGRAPH.CENTER)
    p._p.append(parse_xml(f'<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">{ostr}</m:oMathPara>'))

# TIÊU ĐỀ
add_title("BẢN TÍNH TOÁN THIẾT KẾ ĐỘNG LỰC HỌC VÀ ĐIỀU KHIỂN XE TỰ HÀNH (AGV)")
add_subtitle("HỆ THỐNG THỰC TẾ: XE 1.75 KG, ĐỘNG CƠ GA25 280RPM VÀ DRIVER I2C MKE-M18 BD62130")

# MỤC 1
add_h1("1. TỔNG TẢI TRỌNG VÀ PHÂN BỔ KHỐI LƯỢNG THỰC TẾ")
add_body("Khối lượng thực tế của xe sau khi cân toàn bộ khung gầm, pin, bo mạch điều khiển và 2 cụm động cơ là m_thuc = 1.75 kg. Xe có cấu hình 3 bánh: 1 bánh caster tự lựa dẫn hướng phía trước và 2 bánh chủ động dẫn động độc lập phía sau.")

add_h2("1.1. Công thức tính tổng khối lượng thiết kế có hệ số an toàn")
add_body("Tổng khối lượng thiết kế là khối lượng dùng để tính toán cơ cấu động cơ, bằng khối lượng thực tế nhân với hệ số dự phòng an toàn (chọn k_s = 1.2, tức tính dư thêm 20% tải để dự phòng dây dẫn, ke nối và phụ kiện phát sinh):")
add_eq(r"M = m_{thuc} \cdot k_s")
add_body("Trong đó:")
add_body("M: Tổng khối lượng tính toán của xe (đơn vị: kg).", bold_prefix="• ")
add_body("m_thuc: Khối lượng thực tế cân được của xe (m_thuc = 1.75 kg).", bold_prefix="• ")
add_body("k_s: Hệ số an toàn thiết kế (k_s = 1.2).", bold_prefix="• ")
add_body("Thay số vào công thức:")
add_eq(r"M = 1.75 \cdot 1.2 = 2.10\text{ kg}")

add_h2("1.2. Công thức tính trọng lực toàn xe")
add_body("Trọng lực là lực hút của Trái Đất tác dụng lên toàn bộ khối lượng xe theo phương thẳng đứng hướng vuông góc xuống mặt sàn:")
add_eq(r"P = M \cdot g")
add_body("Trong đó:")
add_body("P: Trọng lực toàn phần của xe (đơn vị: N).", bold_prefix="• ")
add_body("M: Tổng khối lượng tính toán (M = 2.10 kg).", bold_prefix="• ")
add_body("g: Gia tốc trọng trường (lấy chuẩn g = 9.81 m/s²).", bold_prefix="• ")
add_body("Thay số:")
add_eq(r"P = 2.10 \cdot 9.81 = 20.601\text{ N}")

add_h2("1.3. Phân bổ tải trọng lên các cụm bánh xe")
add_body("Để 2 bánh dẫn động phía sau có đủ ma sát bám sàn, tránh hiện tượng trượt quay bánh khi khởi động hoặc phanh, đồng thời bánh caster phía trước đủ nhẹ để chuyển hướng mượt, tải trọng phân bổ theo tỷ lệ 70% dồn vào 2 bánh sau và 30% dồn vào bánh trước:")
add_body("1. Trọng lực dồn lên 2 bánh dẫn động sau:")
add_eq(r"P_{sau} = 0.70 \cdot P = 0.70 \cdot 20.601 = 14.42\text{ N}")
add_body("Mỗi bánh sau chịu tải trọng:")
add_eq(r"P_{banh\_sau} = \dfrac{P_{sau}}{2} = \dfrac{14.42}{2} = 7.21\text{ N} \approx 0.735\text{ kg/bánh}")
add_body("2. Trọng lực dồn lên bánh caster trước:")
add_eq(r"P_{caster} = 0.30 \cdot P = 0.30 \cdot 20.601 = 6.18\text{ N} \approx 0.630\text{ kg}")

# MỤC 2
add_h1("2. THÔNG SỐ HÌNH HỌC BÁNH XE THỰC TẾ")
add_body("Bánh xe đi kèm bộ động cơ GA25 có chu vi vòng tròn ngoài đo được trực tiếp bằng thước dây là C = 23 cm (0.23 m).")

add_h2("2.1. Công thức tính đường kính bánh xe")
add_body("Đường kính bánh xe là chiều dài đoạn thẳng đi qua tâm nối hai điểm đối diện trên đường tròn viền bánh xe:")
add_eq(r"D = \dfrac{C}{\pi}")
add_body("Trong đó:")
add_body("D: Đường kính bánh xe (đơn vị: m).", bold_prefix="• ")
add_body("C: Chu vi bánh xe (C = 0.23 m).", bold_prefix="• ")
add_body("pi: Hằng số Pi (lấy 3.1416).", bold_prefix="• ")
add_body("Thay số:")
add_eq(r"D = \dfrac{0.23}{3.1416} \approx 0.0732\text{ m} = 73.2\text{ mm}")

add_h2("2.2. Công thức tính bán kính bánh xe")
add_body("Bán kính bánh xe là khoảng cách từ trục quay đến mép tiếp xúc mặt sàn:")
add_eq(r"R = \dfrac{D}{2} = \dfrac{0.0732}{2} = 0.0366\text{ m} = 36.6\text{ mm}")

# MỤC 3
add_h1("3. TÍNH TOÁN LỰC CẢN VÀ LỰC KÉO YÊU CẦU")
add_h2("3.1. Công thức tính lực cản lăn")
add_body("Lực cản lăn xuất hiện do ma sát tiếp xúc và biến dạng giữa lốp cao su của bánh xe với mặt sàn phẳng khi xe di chuyển đều:")
add_eq(r"F_{lan} = \mu \cdot P = \mu \cdot M \cdot g")
add_body("Trong đó:")
add_body("F_lan: Lực cản lăn của xe (đơn vị: N).", bold_prefix="• ")
add_body("mu: Hệ số cản lăn sàn gạch/gỗ phẳng (lấy chuẩn mu = 0.02).", bold_prefix="• ")
add_body("P: Trọng lực toàn xe (P = 20.601 N).", bold_prefix="• ")
add_body("Thay số:")
add_eq(r"F_{lan} = 0.02 \cdot 20.601 = 0.412\text{ N}")

add_h2("3.2. Công thức tính lực cản quán tính khi tăng tốc")
add_body("Lực quán tính là lực cản trở gia tốc khi xe bắt đầu tăng tốc từ vị trí đứng yên lên tốc độ định mức theo Định luật II Newton:")
add_eq(r"F_{qt} = M \cdot a")
add_body("Trong đó:")
add_body("F_qt: Lực quán tính khi tăng tốc (đơn vị: N).", bold_prefix="• ")
add_body("M: Tổng khối lượng tính toán (M = 2.10 kg).", bold_prefix="• ")
add_body("a: Gia tốc tăng tốc mong muốn (chọn a = 0.3 m/s² để xe khởi động êm mượt, không giật trượt bánh).", bold_prefix="• ")
add_body("Thay số:")
add_eq(r"F_{qt} = 2.10 \cdot 0.3 = 0.630\text{ N}")

add_h2("3.3. Tổng lực kéo yêu cầu toàn xe và lực đẩy mỗi bánh")
add_body("Tổng lực kéo là hợp lực động cơ cần tạo ra tại thời điểm khởi động tăng tốc:")
add_eq(r"F_{keo} = F_{lan} + F_{qt} = 0.412 + 0.630 = 1.042\text{ N}")
add_body("Mỗi bánh dẫn động sau cần tạo ra một lực đẩy tiếp tuyến:")
add_eq(r"F_{banh} = \dfrac{F_{keo}}{2} = \dfrac{1.042}{2} = 0.521\text{ N}")

# MỤC 4
add_h1("4. ĐỐI CHIẾU THỰC TẾ: TỐC ĐỘ YÊU CẦU VÀ ĐỘNG CƠ 280 RPM ĐÃ MUA")

add_h2("4.1. Tính toán tốc độ vòng quay lý thuyết cần thiết (n_yeucau)")
add_body("Trong không gian phòng thí nghiệm hoặc nhà xưởng có diện tích vừa và nhỏ, tốc độ di chuyển làm việc an toàn và tối ưu của xe tự hành là v = 0.50 m/s (tương đương 1.8 km/h, vận tốc người đi bộ chậm).")
add_body("Công thức tính tốc độ vòng quay yêu cầu của bánh xe:")
add_eq(r"n_{yeucau} = \dfrac{60 \cdot v}{C}")
add_body("Trong đó:")
add_body("n_yeucau: Tốc độ vòng quay lý thuyết cần đạt được trên trục bánh xe (đơn vị: vòng/phút - RPM).", bold_prefix="• ")
add_body("v: Vận tốc di chuyển mong muốn của xe (v = 0.50 m/s).", bold_prefix="• ")
add_body("C: Chu vi vòng tròn bánh xe (C = 0.23 m).", bold_prefix="• ")
add_body("Thay số vào công thức:")
add_eq(r"n_{yeucau} = \dfrac{60 \cdot 0.50}{0.23} \approx 130.4\text{ RPM} \approx 130 - 140\text{ RPM}")
add_body("Như vậy, về mặt lý thuyết thuần túy, xe chỉ cần động cơ có tốc độ quay sau giảm tốc khoảng 130 đến 140 vòng/phút (RPM).")

add_h2("4.2. Thực tế sử dụng động cơ GA25-370 280 RPM và giải pháp điều khiển")
add_body("Thực tế người thiết kế đã mua bộ động cơ GA25-370 12V có tốc độ danh định n_thucte = 280 RPM. Điều này hoàn toàn hợp lý và được giải quyết triệt để trong phần mềm điều khiển như sau:")

add_body("1. Vận tốc tối đa của động cơ 280 RPM khi mở 100% công suất:")
add_eq(r"v_{\max} = \dfrac{n_{thucte} \cdot C}{60} = \dfrac{280 \cdot 0.23}{60} \approx 1.073\text{ m/s} = 3.86\text{ km/h}")

add_body("2. Giải pháp khống chế độ rộng xung PWM trong lập trình:")
add_body("Tỷ lệ phần trăm công suất cần mở để xe đạt đúng tốc độ thiết kế (130 - 140 RPM):")
add_eq(r"\text{Tỷ lệ PWM} = \dfrac{n_{yeucau}}{n_{thucte}} = \dfrac{130}{280} \approx 46.4\% \implies \text{Cài đặt PWM từ 45\% đến 50\%}")
add_body("Trên thang điều khiển PWM 8-bit (từ 0 đến 255):")
add_eq(r"\text{Giá trị PWM cài đặt} = 255 \cdot 48\% \approx 120 - 130")

add_body("3. Ý nghĩa thực tiễn của dải công suất dự phòng (50% còn lại):")
add_body("Việc dùng động cơ 280 RPM và chạy ở mức 45% - 50% PWM mang lại 3 lợi thế thực tế rất lớn:")
add_body("Động cơ hoạt động cực kỳ mát, không bị quá nhiệt, tăng tuổi thọ chổi than và bánh răng hộp số.", bold_prefix="• ")
add_body("Dải công suất còn lại (từ 50% đến 100% PWM) đóng vai trò là 'dự trữ mô-men xoắn': Khi xe leo dốc, vượt thảm dày, vượt gờ sàn hoặc khi bánh caster bị kẹt góc xoay, thuật toán điều khiển PID sẽ tự động bù xung PWM lên 70% - 90% để xe vượt qua chướng ngại vật dễ dàng, sau đó tự hạ về mức 48% khi đường bằng phẳng.", bold_prefix="• ")
add_body("Tránh hoàn toàn hiện tượng xe bị chết máy hoặc tụt áp pin khi gặp tải bất ngờ.", bold_prefix="• ")

# MỤC 5
add_h1("5. TÍNH TOÁN MÔ-MEN XOẮN VÀ ĐỘ BỀN CƠ KHÍ")
add_h2("5.1. Công thức tính mô-men xoắn cần thiết trên trục bánh xe")
add_body("Mô-men xoắn yêu cầu là tích số giữa lực đẩy tiếp tuyến với bán kính bánh xe, chia cho hiệu suất truyền động:")
add_eq(r"T_{banh} = \dfrac{F_{banh} \cdot R}{\eta}")
add_body("Trong đó:")
add_body("T_banh: Mô-men xoắn yêu cầu trên trục bánh (đơn vị: N.m).", bold_prefix="• ")
add_body("F_banh: Lực đẩy mỗi bánh (F_banh = 0.521 N).", bold_prefix="• ")
add_body("R: Bán kính bánh xe (R = 0.0366 m).", bold_prefix="• ")
add_body("eta: Hiệu suất truyền động cơ khí hộp số (chọn eta = 0.85).", bold_prefix="• ")
add_body("Thay số:")
add_eq(r"T_{banh} = \dfrac{0.521 \cdot 0.0366}{0.85} = 0.0224\text{ N}\cdot\text{m} = 0.228\text{ kg}\cdot\text{cm}")
add_body("Áp dụng hệ số an toàn dự trữ công suất k_m = 1.5:")
add_eq(r"T_{chon} \ge T_{banh} \cdot 1.5 = 0.228 \cdot 1.5 = 0.342\text{ kg}\cdot\text{cm}")

add_h2("5.2. Đối chiếu với thông số catalog động cơ GA25-370 280RPM (Hshop)")
add_body("Thông số công bố của nhà sản xuất:")
add_body("Mô-men xoắn định mức: T_dinh_muc = 0.40 kg.cm.", bold_prefix="• ")
add_body("Mô-men xoắn tối đa khi kẹt tải (Stall): T_stall = 2.00 kg.cm.", bold_prefix="• ")
add_body("Dòng không tải: 0.07 A; Dòng khi đầy tải tối đa: 1.8 A.", bold_prefix="• ")
add_body("Đánh giá: Mô-men định mức của GA25 (0.40 kg.cm) lớn hơn mô-men tính toán (0.342 kg.cm), và mô-men cực đại (2.00 kg.cm) gấp gần 6 lần mô-men khởi động. Động cơ hoàn toàn đáp ứng xuất sắc yêu cầu tải của xe.")

# MỤC 6
add_h1("6. TÍNH TOÁN THÔNG SỐ ENCODER VÀ ODOMETRY CHÍNH XÁC")
add_body("Động cơ GA25 tích hợp sẵn Encoder Hall 2 kênh AB lệch pha 90 độ, có 11 xung/vòng/kênh trên đĩa từ. Hộp số giảm tốc có tỷ số truyền K_gear = 21.3 : 1.")

add_h2("6.1. Công thức tính tổng số xung trên 1 vòng quay bánh xe")
add_body("Số xung thực tế sau khi qua hộp số giảm tốc và giải mã 4 cạnh xung (chế độ x4 Quadrature Decoding trên vi điều khiển):")
add_eq(r"N_{rev} = N_{dia} \cdot K_{gear} \cdot 4")
add_body("Trong đó:")
add_body("N_rev: Tổng số xung đếm được khi bánh xe quay trọn 1 vòng (đơn vị: xung/vòng).", bold_prefix="• ")
add_body("N_dia: Số xung đĩa từ trên 1 kênh (N_dia = 11 xung/vòng).", bold_prefix="• ")
add_body("K_gear: Tỷ số truyền hộp số (K_gear = 21.3).", bold_prefix="• ")
add_body("4: Hệ số giải mã Quadrature x4 (đếm cả 2 sườn lên và xuống của kênh A và B).", bold_prefix="• ")
add_body("Thay số:")
add_eq(r"N_{rev} = 11 \cdot 21.3 \cdot 4 = 937.2\text{ xung/vòng}")

add_h2("6.2. Công thức tính độ phân giải quãng đường di chuyển")
add_body("Độ phân giải quãng đường là cự ly xe lăn bánh tương ứng với 1 xung Encoder đọc được:")
add_eq(r"\Delta s_{pulse} = \dfrac{C}{N_{rev}} = \dfrac{0.23}{937.2} \approx 0.000245\text{ m} = 0.245\text{ mm/xung}")
add_body("Nhận xét: Mỗi xung tương đương 0.245 mm, cho phép vi điều khiển tính toán vị trí xe với độ mịn cực cao, hạn chế tích lũy sai số trôi trong thuật toán Odometry.")

# MỤC 7
add_h1("7. TÍNH TOÁN ĐỘ TƯƠNG THÍCH CỦA DRIVER MKE-M18 BD62130 I2C")
add_body("Hệ thống sử dụng module điều khiển động cơ DC MKE-M18 BD62130 giao tiếp qua I2C của MakerEdu.")

add_h2("7.1. Khả năng tương thích dòng điện và điện áp")
add_body("1. Điện áp cấp động cơ (Vin): Module hỗ trợ 8.0V - 13.0VDC, hoàn toàn tương thích với động cơ 12VDC và các loại pin Li-ion 3S (11.1V - 12.6V).")
add_body("2. Dòng điện chịu tải: Module chịu tải liên tục 3.0A mỗi kênh và dòng đỉnh 4.0A. Động cơ GA25 chạy có tải bình thường chỉ ăn dòng 0.2A - 0.4A, và dòng cực đại khi kẹt tải là 1.8A. Driver dư tải an toàn, hoạt động mát và không bị sụt áp.")

add_h2("7.2. Lợi thế kiến trúc I2C trong lập trình xe tự hành")
add_body("1. Tiết kiệm chân vi điều khiển: Module chỉ dùng đúng 2 chân I2C (SDA, SCL) với địa chỉ mặc định 0x40. Toàn bộ các chân GPIO ngắt ngoài trên Arduino/ESP32/STM32 được giữ lại để đọc tín hiệu 2 bộ Encoder và cảm biến.")
add_body("2. Giảm tải vi điều khiển chính: Mạch tích hợp chip 32-bit ARM Cortex-M0+ phụ tự động tạo xung PWM độc lập (tần số 500Hz - 20kHz). Vi điều khiển chính chỉ cần gửi lệnh đặt tốc độ I2C trong dải 120 - 130 (ứng với 48% công suất), rảnh tay 100% để tính toán Odometry và thuật toán dẫn đường.")

doc.save(out_docx)
print("Updated successfully to single file:", out_docx)
