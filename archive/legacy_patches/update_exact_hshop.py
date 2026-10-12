# -*- coding: utf-8 -*-
"""
Cập nhật tài liệu tính toán thiết kế xe tự hành AGV:
- Khối lượng cân thực tế: 1.75 kg (tính toán an toàn 2.1 kg)
- Bộ động cơ chính xác từ Hshop: GA25-370 DC Servo 12V 280RPM
- Tỷ số truyền: 21.3:1
- Encoder Hall 2 kênh AB: 11 xung/vòng/kênh (44 xung/vòng ở chế độ x4)
- Bánh xe chu vi 23 cm (D = 73.2 mm)
File lưu tại: D:\\HoMinhKhang\\PROJECT\\GITHUB_PROJECT\\CODE XE TU HANH\\TINH_TOAN_THIET_KE_XE_TU_HANH_FINAL.docx
"""
import os
import lxml.etree as ET
import latex2mathml.converter
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml

target_dir = r"D:\HoMinhKhang\PROJECT\GITHUB_PROJECT\CODE XE TU HANH"
os.makedirs(target_dir, exist_ok=True)
out_docx = os.path.join(target_dir, "TINH_TOAN_THIET_KE_XE_TU_HANH_FINAL.docx")

xslt_path = r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"
xslt = ET.parse(xslt_path)
xslt_transform = ET.XSLT(xslt)

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(0.8)

BLACK = RGBColor(0, 0, 0)
PRIMARY = RGBColor(0, 51, 102)

def format_p(p, ls=1.3, sa=4, sb=0, al=WD_ALIGN_PARAGRAPH.LEFT, indent=0):
    p.paragraph_format.line_spacing = ls
    p.paragraph_format.space_after = Pt(sa)
    p.paragraph_format.space_before = Pt(sb)
    p.alignment = al
    if indent > 0:
        p.paragraph_format.left_indent = Inches(indent)
    return p

def add_title(text):
    p = doc.add_paragraph()
    format_p(p, 1.2, 4, 10, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = BLACK

def add_subtitle(text):
    p = doc.add_paragraph()
    format_p(p, 1.2, 14, 2, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = PRIMARY

def add_h1(text):
    p = doc.add_paragraph()
    format_p(p, 1.2, 4, 12, WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = PRIMARY

def add_h2(text):
    p = doc.add_paragraph()
    format_p(p, 1.2, 3, 8, WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = BLACK

def add_body(text, bold_prefix="", indent=0):
    p = doc.add_paragraph()
    format_p(p, 1.3, 3, 0, WD_ALIGN_PARAGRAPH.LEFT, indent)
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.font.name = "Times New Roman"
        rb.font.size = Pt(12)
        rb.font.bold = True
        rb.font.color.rgb = BLACK
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.color.rgb = BLACK

def add_eq(latex_code):
    mathml = latex2mathml.converter.convert(latex_code)
    tree = ET.fromstring(mathml)
    omml = xslt_transform(tree)
    ostr = ET.tostring(omml, encoding="utf-8").decode("utf-8")
    p = doc.add_paragraph()
    format_p(p, 1.15, 4, 4, WD_ALIGN_PARAGRAPH.CENTER)
    p._p.append(parse_xml(f'<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">{ostr}</m:oMathPara>'))

add_title("BẢN TÍNH TOÁN THIẾT KẾ ĐỘNG LỰC HỌC VÀ ĐIỀU KHIỂN XE TỰ HÀNH (AGV)")
add_subtitle("CẤU HÌNH XE 1.75 KG (HỆ SỐ AN TOÀN 2.1 KG) - ĐỘNG CƠ GA25-370 12V 280RPM CÓ ENCODER")

# PHẦN 1
add_h1("1. THÔNG SỐ CẤU HÌNH VÀ TỔNG TẢI TRỌNG THIẾT KẾ")
add_body("Khối lượng thực tế xe cân được là 1.75 kg (bao gồm toàn bộ khung gầm, pin, mạch điều khiển và 2 cụm động cơ). Cấu hình di chuyển gồm 1 bánh caster tự lựa dẫn hướng phía trước và 2 bánh dẫn động chủ động phía sau.")

add_h2("1.1. Công thức tính tổng khối lượng có hệ số an toàn")
add_body("Tổng khối lượng thiết kế là khối lượng dùng để tính toán động cơ, bằng khối lượng thực tế nhân với hệ số dự phòng an toàn (k_s = 1.2, dự phòng thêm 20% tải):")
add_eq(r"M = m_{thuc} \cdot k_s")
add_body("Trong đó:")
add_body("M: Tổng khối lượng tính toán (đơn vị: kg).", bold_prefix="• ")
add_body("m_thuc: Khối lượng thực tế cân được của xe (m_thuc = 1.75 kg).", bold_prefix="• ")
add_body("k_s: Hệ số an toàn thiết kế (chọn k_s = 1.2).", bold_prefix="• ")
add_body("Thay số vào công thức:")
add_eq(r"M = 1.75 \cdot 1.2 = 2.10\text{ kg}")

add_h2("1.2. Công thức tính trọng lực toàn xe")
add_body("Trọng lực là lực hút của Trái Đất tác dụng lên toàn bộ khối lượng xe theo phương thẳng đứng:")
add_eq(r"P = M \cdot g")
add_body("Trong đó:")
add_body("P: Trọng lực toàn phần của xe (đơn vị: N).", bold_prefix="• ")
add_body("M: Tổng khối lượng tính toán (M = 2.10 kg).", bold_prefix="• ")
add_body("g: Gia tốc trọng trường (g = 9.81 m/s²).", bold_prefix="• ")
add_body("Thay số vào công thức:")
add_eq(r"P = 2.10 \cdot 9.81 = 20.601\text{ N}")

add_h2("1.3. Phân bổ tải trọng lên các cụm bánh xe")
add_body("Để 2 bánh sau bám sàn tốt không trượt và bánh caster trước xoay chuyển hướng mượt, tải trọng phân bổ 70% đè lên 2 bánh sau và 30% đè lên bánh trước:")
add_body("1. Trọng lực đè lên 2 bánh dẫn động sau:")
add_eq(r"P_{sau} = 0.70 \cdot P = 0.70 \cdot 20.601 = 14.42\text{ N}")
add_body("Mỗi bánh sau chịu tải trọng:")
add_eq(r"P_{banh\_sau} = \dfrac{P_{sau}}{2} = \dfrac{14.42}{2} = 7.21\text{ N} \approx 0.735\text{ kg/bánh}")

add_body("2. Trọng lực đè lên 1 bánh caster trước:")
add_eq(r"P_{caster} = 0.30 \cdot P = 0.30 \cdot 20.601 = 6.18\text{ N} \approx 0.630\text{ kg}")

# PHẦN 2
add_h1("2. THÔNG SỐ HÌNH HỌC BÁNH XE")
add_body("Bánh xe đi kèm bộ GA25 có chu vi vòng tròn ngoài đo được C = 23 cm (0.23 m).")

add_h2("2.1. Công thức tính đường kính bánh xe")
add_body("Đường kính bánh xe được tính từ chu vi theo công thức hình học:")
add_eq(r"D = \dfrac{C}{\pi}")
add_body("Trong đó:")
add_body("D: Đường kính bánh xe (đơn vị: m).", bold_prefix="• ")
add_body("C: Chu vi bánh xe (C = 0.23 m).", bold_prefix="• ")
add_body("pi: Hằng số Pi (lấy 3.1416).", bold_prefix="• ")
add_body("Thay số:")
add_eq(r"D = \dfrac{0.23}{3.1416} \approx 0.0732\text{ m} = 73.2\text{ mm}")

add_h2("2.2. Công thức tính bán kính bánh xe")
add_body("Bán kính bánh xe bằng một nửa đường kính:")
add_eq(r"R = \dfrac{D}{2} = \dfrac{0.0732}{2} = 0.0366\text{ m} = 36.6\text{ mm}")

# PHẦN 3
add_h1("3. TÍNH TOÁN CÁC THÀNH PHẦN LỰC CẢN CHUYỂN ĐỘNG")
add_h2("3.1. Công thức tính lực cản lăn")
add_body("Lực cản lăn sinh ra do ma sát tiếp xúc giữa bánh xe cao su và mặt sàn cứng:")
add_eq(r"F_{lan} = \mu \cdot P = \mu \cdot M \cdot g")
add_body("Trong đó:")
add_body("F_lan: Lực cản lăn (đơn vị: N).", bold_prefix="• ")
add_body("mu: Hệ số cản lăn sàn gạch/gỗ phẳng (mu = 0.02).", bold_prefix="• ")
add_body("P: Trọng lực toàn xe (P = 20.601 N).", bold_prefix="• ")
add_body("Thay số:")
add_eq(r"F_{lan} = 0.02 \cdot 20.601 = 0.412\text{ N}")

add_h2("3.2. Công thức tính lực quán tính khi tăng tốc")
add_body("Lực quán tính là lực cản trở khi xe tăng tốc từ 0 lên tốc độ làm việc theo Định luật II Newton:")
add_eq(r"F_{qt} = M \cdot a")
add_body("Trong đó:")
add_body("F_qt: Lực quán tính (đơn vị: N).", bold_prefix="• ")
add_body("M: Tổng khối lượng tính toán (M = 2.10 kg).", bold_prefix="• ")
add_body("a: Gia tốc tăng tốc của xe (chọn a = 0.3 m/s²).", bold_prefix="• ")
add_body("Thay số:")
add_eq(r"F_{qt} = 2.10 \cdot 0.3 = 0.630\text{ N}")

add_h2("3.3. Tổng lực kéo yêu cầu toàn xe")
add_body("Tổng lực kéo cần thiết tại thời điểm khởi động tăng tốc:")
add_eq(r"F_{keo} = F_{lan} + F_{qt} = 0.412 + 0.630 = 1.042\text{ N}")
add_body("Mỗi bánh dẫn động sau chỉ cần tạo ra một lực đẩy:")
add_eq(r"F_{banh} = \dfrac{F_{keo}}{2} = \dfrac{1.042}{2} = 0.521\text{ N}")

# PHẦN 4
add_h1("4. TÍNH TOÁN MÔ-MEN XOẮN VÀ KIỂM CHỨNG ĐỘNG CƠ GA25-370 280RPM")
add_h2("4.1. Công thức tính mô-men xoắn cần thiết trên trục bánh")
add_body("Mô-men xoắn là tích của lực đẩy tiếp tuyến với bán kính bánh xe, chia cho hiệu suất truyền động:")
add_eq(r"T_{banh} = \dfrac{F_{banh} \cdot R}{\eta}")
add_body("Trong đó:")
add_body("T_banh: Mô-men xoắn yêu cầu trên trục bánh (đơn vị: N.m).", bold_prefix="• ")
add_body("F_banh: Lực đẩy mỗi bánh (F_banh = 0.521 N).", bold_prefix="• ")
add_body("R: Bán kính bánh xe (R = 0.0366 m).", bold_prefix="• ")
add_body("eta: Hiệu suất truyền động (eta = 0.85).", bold_prefix="• ")
add_body("Thay số:")
add_eq(r"T_{banh} = \dfrac{0.521 \cdot 0.0366}{0.85} = 0.0224\text{ N}\cdot\text{m}")
add_body("Quy đổi sang kg.cm:")
add_eq(r"T_{banh} = 0.0224 \cdot 10.197 = 0.228\text{ kg}\cdot\text{cm}")
add_body("Mô-men xoắn yêu cầu có hệ số an toàn dự phòng k_m = 1.5:")
add_eq(r"T_{chon} \ge 0.228 \cdot 1.5 = 0.342\text{ kg}\cdot\text{cm}")

add_h2("4.2. Đối chiếu trực tiếp với Catalog động cơ GA25-370 12V 280RPM (Hshop)")
add_body("Thông số kỹ thuật chính thức từ nhà sản xuất:")
add_body("Điện áp danh định: 12 VDC.", bold_prefix="• ")
add_body("Tỷ số truyền hộp số kim loại: 21.3 : 1.", bold_prefix="• ")
add_body("Tốc độ không tải: 280 RPM; Tốc độ có tải trung bình: 215 RPM.", bold_prefix="• ")
add_body("Mô-men xoắn định mức của động cơ: 0.40 kg.cm.", bold_prefix="• ")
add_body("Mô-men xoắn tối đa (Stall/Max): 2.00 kg.cm.", bold_prefix="• ")
add_body("Dòng không tải: 0.07 A; Dòng tối đa khi đầy tải: 1.8 A.", bold_prefix="• ")

add_body("Đánh giá: Mô-men định mức của động cơ đạt 0.40 kg.cm, lớn hơn mô-men yêu cầu tính toán (0.342 kg.cm), và mô-men tối đa đạt 2.0 kg.cm (gấp gần 6 lần tải khởi động). Động cơ hoàn toàn đáp ứng xuất sắc, hoạt động mát và bền bỉ.")

# PHẦN 5
add_h1("5. TÍNH TOÁN THÔNG SỐ ENCODER VÀ ĐỊNH VỊ ODOMETRY CHÍNH XÁC")
add_body("Động cơ tích hợp sẵn Encoder Hall 2 kênh AB lệch pha 90 độ.")

add_h2("5.1. Công thức tính số xung Encoder trên 1 vòng quay bánh xe")
add_body("Số xung thực tế sau khi qua hộp số giảm tốc và giải mã 4 cạnh xung (chế độ x4 Quadrature Decoding):")
add_eq(r"N_{rev} = N_{dia} \cdot K_{gear} \cdot 4")
add_body("Trong đó:")
add_body("N_rev: Tổng số xung vi điều khiển đếm được khi bánh xe quay trọn 1 vòng (xung/vòng).", bold_prefix="• ")
add_body("N_dia: Số xung đĩa từ trên 1 kênh của động cơ (N_dia = 11 xung/vòng).", bold_prefix="• ")
add_body("K_gear: Tỷ số truyền của hộp số giảm tốc GA25 (K_gear = 21.3).", bold_prefix="• ")
add_body("4: Hệ số giải mã Quadrature x4 (đếm cả 2 sườn lên và xuống của 2 kênh A và B).", bold_prefix="• ")
add_body("Thay số:")
add_eq(r"N_{rev} = 11 \cdot 21.3 \cdot 4 = 937.2\text{ xung/vòng}")

add_h2("5.2. Công thức tính độ phân giải quãng đường (Khoảng cách trên mỗi xung)")
add_body("Độ phân giải quãng đường là cự ly xe di chuyển ứng với 1 xung Encoder vi điều khiển đọc được:")
add_eq(r"\Delta s_{pulse} = \dfrac{C}{N_{rev}} = \dfrac{0.23}{937.2} \approx 0.000245\text{ m} = 0.245\text{ mm/xung}")
add_body("Nhận xét: Mỗi xung ứng với chỉ 0.245 mm. Đây là độ phân giải cực kỳ cao, giúp thuật toán Odometry định vị vị trí xe cực kỳ mịn và chính xác.")


# PHẦN 6
add_h1("6. TÍNH TOÁN VÀ ĐÁNH GIÁ DRIVER ĐIỀU KHIỂN ĐỘNG CƠ (MKE-M18 BD62130)")
add_body("Hệ thống sử dụng module điều khiển động cơ DC MKE-M18 tích hợp cầu H kép BD62130 giao tiếp qua I2C của MakerEdu.")

add_h2("6.1. Kiểm tra độ tương thích dòng điện và điện áp")
add_body("1. Điện áp cấp động cơ (Vin): Module hỗ trợ dải 8.0V - 13.0VDC, hoàn toàn tương thích hoàn hảo với động cơ GA25 12VDC và các khối pin 3S (11.1V - 12.6V).")
add_body("2. Dòng điện làm việc liên tục: Module chịu được 3.0A mỗi kênh. Động cơ GA25 chạy có tải bình thường chỉ tiêu thụ 0.2A - 0.4A, và dòng cực đại khi kẹt tải (stall) là 1.8A. Do đó dòng định mức 3.0A của Driver dư sức gánh tải mà không lo bị nóng hay sụt áp.")
add_body("3. Dòng đỉnh tức thời: Module chịu được dòng đỉnh lên đến 4.0A mỗi kênh, đảm bảo an toàn tuyệt đối khi xe khởi động hoặc đảo chiều đột ngột.")

add_h2("6.2. Ưu điểm kiến trúc giao tiếp I2C đối với hệ thống xe tự hành")
add_body("1. Tiết kiệm chân vi điều khiển tối đa: Thay vì phải dùng 4 đến 6 chân GPIO và Timer PWM trên Arduino/ESP32/STM32, module MKE-M18 chỉ chiếm đúng 2 chân I2C (SDA và SCL). Toàn bộ các chân GPIO còn lại có thể dành trọn vẹn cho việc đọc xung Encoder ngắt ngoài (External Interrupt) và gắn cảm biến LiDAR/siêu âm.")
add_body("2. Giảm tải xử lý (Hardware Offloading): Module tích hợp sẵn vi điều khiển phụ 32-bit ARM Cortex-M0+ đảm nhiệm việc băm xung PWM độc lập (tần số 500Hz - 20kHz) và đảo chiều cầu H, giúp CPU chính tập trung 100% tài nguyên cho giải thuật Odometry và bám quỹ đạo.")
add_body("3. Khống chế dải PWM an toàn: Động cơ 280 RPM có dải tốc độ làm việc tối ưu trong nhà ở mức 40% - 55% PWM. Trên thanh ghi I2C điều khiển (giá trị 0 - 255), vi điều khiển chỉ cần gửi lệnh đặt tốc độ trong khoảng PWM = 100 đến 140.")


doc.save(out_docx)
print("File updated with full Hshop specs:", out_docx)
